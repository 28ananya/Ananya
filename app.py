from flask import Flask, render_template, request, send_file
import os
import re
import xlwt
import docx2txt
from PyPDF2 import PdfReader
from docx import Document
import zipfile
import os
from comtypes import client
from docx import Document



app = Flask(__name__)


def convert_doc_to_pdf(doc_file, pdf_file):
    word = client.CreateObject('Word.Application')
    doc = word.Documents.Open(doc_file)
    doc.SaveAs(pdf_file, FileFormat=17)  # FileFormat 17 is for PDF
    doc.Close()
    word.Quit()

def extract_information_from_cv(cv_text):
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    phone_pattern = r'\b(?:\+?(\d{1,3}))?[-. (]*(\d{3})[-. )]*(\d{3})[-. ]*(\d{4})\b'
    emails = re.findall(email_pattern, cv_text)
    phones = re.findall(phone_pattern, cv_text)
    formatted_phones = ['({}) {}-{}'.format(phone[1], phone[2], phone[3]) for phone in phones]
    return emails, formatted_phones, cv_text

def process_docx(docx_file):
    return docx2txt.process(docx_file)

def process_doc(doc_file):
    if isinstance(doc_file, io.BytesIO):
        temp_docx_file = "temp.docx"
        with open(temp_docx_file, "wb") as temp_file:
            temp_file.write(doc_file.read())
        doc = Document(temp_docx_file)
        text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
        os.remove(temp_docx_file)
        return text
    else:
        return "error"
    


def process_pdf(pdf_file):
    reader = PdfReader(pdf_file)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text

import io

def extract_text_from_file(file_name, file_content):
    _, file_extension = os.path.splitext(file_name)
    if file_extension.lower() == '.docx':
        return process_docx(file_content)
    elif file_extension.lower() == '.doc':
        return process_doc(file_content)
    elif file_extension.lower() == '.pdf':
        return process_pdf(file_content)
    else:
        return ""

    



def process_zip_file(zip_file_path):
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet('Resume Information')
    sheet.write(0, 0, 'File Name')
    sheet.write(0, 1, 'Email Addresses')
    sheet.write(0, 2, 'Contact Numbers')
    sheet.write(0, 3, 'Overall Text')
    row_index = 1
    with zipfile.ZipFile(zip_file_path, 'r') as zip_ref:
        for file_info in zip_ref.infolist():
            if not isinstance(file_info, zipfile.ZipInfo):
                continue  
            file_name = file_info.filename
            _, file_extension = os.path.splitext(file_name)
            if file_extension.lower() in ['.docx', '.doc', '.pdf']:
                with zip_ref.open(file_name) as file:
                    cv_text = extract_text_from_file(file_name, file)
                    if cv_text:
                        emails, phones, overall_text = extract_information_from_cv(cv_text)
                        sheet.write(row_index, 0, file_name)
                        sheet.write(row_index, 1, ', '.join(emails))
                        sheet.write(row_index, 2, ', '.join(phones))
                        sheet.write(row_index, 3, overall_text)
                        row_index += 1
    return workbook


@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload():
    uploaded_file = request.files['file']
    if uploaded_file.filename != '':
        file_path = os.path.join('', uploaded_file.filename)
        uploaded_file.save(file_path)
        workbook = process_zip_file(file_path)
        excel_file_path = os.path.join('', 'resume_information7.xls')
        workbook.save(excel_file_path)
        return send_file(excel_file_path, as_attachment=True)
    return 'No file uploaded.'

if __name__ == '__main__':
    app.run(debug=True)
